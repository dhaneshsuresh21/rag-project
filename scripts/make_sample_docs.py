"""Generate synthetic sample solution docs (for local development only).
Usage: python scripts/make_sample_docs.py  -> data/docs/SAMPLE-*.docx
Replace them with real solution docs when available."""
from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parents[1] / "data" / "docs"


def _table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            t.cell(i, j).text = val


def refunds_doc(path: Path):
    d = Document()
    d.add_paragraph("GAP-1423 Multi-currency Refunds", style="Title")
    d.add_paragraph("Client: General")
    d.add_paragraph("Version: 1.0")

    d.add_heading("1 Overview", level=1)
    d.add_paragraph(
        "This solution enables refunds in the original payment currency instead of converting all refunds to the "
        "base currency. It applies to card and bank-transfer payments created after the feature is enabled.")

    d.add_heading("2 Business requirement", level=1)
    d.add_paragraph(
        "Customers paying in EUR, GBP or USD must receive refunds in the same currency to avoid exchange-rate "
        "losses. Finance teams need per-currency refund limits and an approval step for large refunds.")

    d.add_heading("3 Configuration", level=1)
    d.add_heading("3.1 Enable multi-currency refunds", level=2)
    for step in ["Go to Settings > Payments > Refund Rules.",
                 "Switch on the Multi-currency Refunds toggle.",
                 "Click Save. The change takes effect immediately for new refunds."]:
        d.add_paragraph(step, style="List Number")

    d.add_heading("3.2 Refund limits per currency", level=2)
    d.add_paragraph(
        "Refund limits are configured in the Currency Limits table on the Refund Rules screen. Each row defines a "
        "currency, the maximum refund amount and whether approval is required above that amount.")
    _table(d, [
        ["Currency", "Max Refund Amount", "Approval Required", "Approver Role"],
        ["EUR", "5,000", "Yes", "Finance Manager"],
        ["GBP", "4,000", "Yes", "Finance Manager"],
        ["USD", "6,000", "No", "-"],
    ])
    d.add_paragraph(
        "Limits apply only to refunds created after the rule is saved. Pending refunds keep the previous limit. "
        "The field refund_limit_ccy stores the ISO 4217 currency code of each limit.")

    d.add_heading("3.3 Approval workflow", level=2)
    d.add_paragraph(
        "When a refund exceeds the limit and Approval Required is Yes, the refund moves to status PENDING_APPROVAL. "
        "The user with the Approver Role receives an email and an in-app notification. If the approver rejects the "
        "refund, its status becomes REJECTED and the requester is notified. Approvals expire after 5 business days.")

    d.add_heading("4 Error codes", level=1)
    _table(d, [
        ["Code", "Meaning", "Resolution"],
        ["RF-409", "Refund currency differs from payment currency", "Enable Multi-currency Refunds (section 3.1)"],
        ["RF-422", "Refund amount above limit without approver", "Assign an Approver Role in Currency Limits"],
    ])

    d.add_heading("5 Known limitations", level=1)
    d.add_paragraph("Partial refunds in a different currency than the original payment are not supported.")
    d.add_paragraph("Cryptocurrency payments are out of scope for this solution.")
    d.save(path)


def reconciliation_doc(path: Path):
    d = Document()
    d.add_paragraph("FEAT-2051 Automated Invoice Reconciliation", style="Title")
    d.add_paragraph("Client: General")

    d.add_heading("1 Overview", level=1)
    d.add_paragraph(
        "Automated Invoice Reconciliation matches incoming bank statement lines to open invoices every night and "
        "flags exceptions for manual review.")

    d.add_heading("2 Matching rules", level=1)
    d.add_paragraph(
        "A statement line is matched to an invoice when the invoice number appears in the payment reference and the "
        "amount is within the tolerance. The default tolerance is 0.50 in the invoice currency and can be changed in "
        "Settings > Reconciliation > Tolerance.")
    d.add_paragraph(
        "If several invoices match, the oldest open invoice is selected. Lines without a match go to the "
        "Exceptions queue.")

    d.add_heading("3 Scheduling", level=1)
    d.add_paragraph(
        "The reconciliation job runs daily at 02:00 server time. Administrators can trigger a manual run from "
        "Reconciliation > Run now. Only one run can be active at a time.")

    d.add_heading("4 Exceptions queue", level=1)
    for step in ["Open Reconciliation > Exceptions.",
                 "Select a statement line and click Find invoice.",
                 "Choose the invoice and click Match, or click Ignore to dismiss the line."]:
        d.add_paragraph(step, style="List Number")
    d.add_paragraph("Ignored lines are kept for 90 days and can be restored from the Ignored tab.")
    d.save(path)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    refunds_doc(OUT / "SAMPLE-GAP-1423-multi-currency-refunds.docx")
    reconciliation_doc(OUT / "SAMPLE-FEAT-2051-invoice-reconciliation.docx")
    print(f"Sample docs written to {OUT}")
