from docx import Document

from app.config import ChunkingCfg, ParsingCfg
from app.ingestion.chunker import chunk_document, embed_text
from app.ingestion.docx_parser import parse_docx


def _make_doc(path, styled=True):
    d = Document()
    d.add_paragraph("GAP-77 Test Feature", style="Title")
    d.add_paragraph("Client: ACME Corp")
    if styled:
        d.add_heading("1 Overview", level=1)
    else:
        d.add_paragraph("1 Overview")
    d.add_paragraph("Short intro.")
    if styled:
        d.add_heading("2 Configuration", level=1)
        d.add_heading("2.1 Limits", level=2)
    else:
        d.add_paragraph("2 Configuration")
        d.add_paragraph("2.1 Limits")
    d.add_paragraph("Long text sentence about limits. " * 200)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Field", "Value"
    t.cell(1, 0).text, t.cell(1, 1).text = "max|limit", "5"
    d.save(path)
    return path


def test_parse_styled_doc(tmp_path):
    doc = parse_docx(_make_doc(tmp_path / "GAP-77 test.docx"), ParsingCfg())
    assert doc.title == "GAP-77 Test Feature"
    assert doc.feature_id == "GAP-77"
    assert doc.doc_type == "gap"
    assert doc.client == "ACME Corp"
    headings = [(b.level, b.text) for b in doc.blocks if b.kind == "heading"]
    assert headings == [(1, "1 Overview"), (1, "2 Configuration"), (2, "2.1 Limits")]
    table = next(b for b in doc.blocks if b.kind == "table")
    assert "| Field | Value |" in table.text and "max\\|limit" in table.text
    assert not doc.warnings


def test_parse_fallback_headings(tmp_path):
    doc = parse_docx(_make_doc(tmp_path / "plain.docx", styled=False), ParsingCfg())
    assert doc.warnings
    assert [(b.level, b.text) for b in doc.blocks if b.kind == "heading"] == [
        (1, "1 Overview"), (1, "2 Configuration"), (2, "2.1 Limits")]


def test_chunking_splits_long_sections_with_overlap(tmp_path):
    cfg = ChunkingCfg(max_tokens=800, target_tokens=650, overlap_tokens=100, min_tokens=60)
    doc = parse_docx(_make_doc(tmp_path / "GAP-77.docx"), ParsingCfg())
    chunks = chunk_document(doc, cfg)
    limits = [c for c in chunks if c["section_path"] == "2 Configuration > 2.1 Limits"]
    assert len(limits) >= 2
    assert all(c["tokens"] <= cfg.max_tokens + cfg.overlap_tokens for c in limits)
    assert limits[0]["parts"] == len(limits)
    assert limits[1]["text"].startswith("...")  # overlap carried from the previous part
    # tiny pre-heading text ("Client: ...") and tiny "1 Overview" are top-level siblings -> merged
    assert chunks[0]["section_path"] == "Overview / 1 Overview" and "Short intro." in chunks[0]["text"]
    assert embed_text(chunks[0]).startswith("GAP-77 Test Feature | GAP-77 | ")


def test_tiny_siblings_merge_and_placeholders_drop(tmp_path):
    d = Document()
    d.add_heading("Extracts", level=1)
    d.add_paragraph("The CSV file specification is below:")
    d.add_heading("Header Record", level=2)
    d.add_paragraph("The header record contains the file creation date and the sender code. " * 10)
    d.add_heading("Fire Job", level=1)
    d.add_heading("Name of Fire Job", level=2)
    d.add_paragraph("Generate report.")
    d.add_heading("Fire Job Id", level=2)
    d.add_paragraph("GEN_JV")
    d.add_heading("Webservices", level=1)
    d.add_paragraph("NA")
    d.add_heading("Impact on Clients", level=1)
    d.add_paragraph("ALL")
    d.save(tmp_path / "PRAC-123_test_v1.docx")
    doc = parse_docx(tmp_path / "PRAC-123_test_v1.docx", ParsingCfg())
    assert doc.feature_id == "PRAC-123"
    chunks = chunk_document(doc, ChunkingCfg())
    paths = [c["section_path"] for c in chunks]
    assert paths == ["Extracts > Header Record", "Fire Job > Name of Fire Job / Fire Job Id", "Impact on Clients"]
    assert chunks[0]["text"].startswith("Extracts:\nThe CSV file specification is below:")
    assert "Fire Job Id:\nGEN_JV" in chunks[1]["text"]
